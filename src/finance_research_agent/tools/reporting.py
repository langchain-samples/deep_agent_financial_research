"""Tools for generating research reports in multiple formats."""

import json
import os
from typing import Dict, Any, Optional
from datetime import datetime
from pathlib import Path
from langchain.tools import tool
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

from ..utils.validators import validate_output_format
from ..utils.formatters import format_currency, format_percentage, create_report_metadata


@tool
def generate_excel_report(data: Dict[str, Any], output_path: str) -> Dict[str, Any]:
    """Generate an Excel report from research data.
    
    Args:
        data: Dictionary containing research data
        output_path: Path where the Excel file should be saved
    
    Returns:
        Dictionary with status and file path
    """
    try:
        # Ensure output directory exists
        output_dir = os.path.dirname(output_path)
        if output_dir:
            os.makedirs(output_dir, exist_ok=True)
        
        # Create workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Research Report"
        
        # Header style
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        
        row = 1
        
        # Add metadata
        if "metadata" in data:
            metadata = data["metadata"]
            ws.cell(row=row, column=1, value="Ticker:").font = Font(bold=True)
            ws.cell(row=row, column=2, value=metadata.get("ticker", "N/A"))
            row += 1
            ws.cell(row=row, column=1, value="Generated At:").font = Font(bold=True)
            ws.cell(row=row, column=2, value=metadata.get("generated_at", "N/A"))
            row += 2
        
        # Add stock info
        if "stock_info" in data:
            row += 1
            ws.cell(row=row, column=1, value="Stock Information").font = header_font
            ws.cell(row=row, column=1).fill = header_fill
            row += 1
            
            stock_info = data["stock_info"]
            for key, value in stock_info.items():
                ws.cell(row=row, column=1, value=key.replace("_", " ").title()).font = Font(bold=True)
                ws.cell(row=row, column=2, value=str(value))
                row += 1
        
        # Add metrics
        if "metrics" in data:
            row += 1
            ws.cell(row=row, column=1, value="Financial Metrics").font = header_font
            ws.cell(row=row, column=1).fill = header_fill
            row += 1
            
            metrics = data["metrics"]
            for key, value in metrics.items():
                ws.cell(row=row, column=1, value=key.replace("_", " ").title()).font = Font(bold=True)
                ws.cell(row=row, column=2, value=str(value))
                row += 1
        
        # Add analysis
        if "analysis" in data:
            row += 1
            ws.cell(row=row, column=1, value="Analysis").font = header_font
            ws.cell(row=row, column=1).fill = header_fill
            row += 1
            
            analysis = data["analysis"]
            if isinstance(analysis, dict):
                for key, value in analysis.items():
                    ws.cell(row=row, column=1, value=key.replace("_", " ").title()).font = Font(bold=True)
                    ws.cell(row=row, column=2, value=str(value))
                    row += 1
        
        # Save workbook
        wb.save(output_path)
        
        return {
            "status": "success",
            "format": "excel",
            "output_path": output_path,
            "file_size": os.path.getsize(output_path) if os.path.exists(output_path) else 0,
        }
    except Exception as e:
        return {"error": f"Failed to generate Excel report: {str(e)}"}


@tool
def generate_pdf_report(data: Dict[str, Any], output_path: str) -> Dict[str, Any]:
    """Generate a PDF report from research data.
    
    Args:
        data: Dictionary containing research data
        output_path: Path where the PDF file should be saved
    
    Returns:
        Dictionary with status and file path
    """
    try:
        # Ensure output directory exists
        output_dir = os.path.dirname(output_path)
        if output_dir:
            os.makedirs(output_dir, exist_ok=True)
        
        # Create PDF document
        doc = SimpleDocTemplate(output_path, pagesize=letter)
        story = []
        styles = getSampleStyleSheet()
        
        # Title
        title = Paragraph("Financial Research Report", styles["Title"])
        story.append(title)
        story.append(Spacer(1, 12))
        
        # Metadata
        if "metadata" in data:
            metadata = data["metadata"]
            meta_text = f"Ticker: {metadata.get('ticker', 'N/A')}<br/>"
            meta_text += f"Generated: {metadata.get('generated_at', 'N/A')}"
            meta_para = Paragraph(meta_text, styles["Normal"])
            story.append(meta_para)
            story.append(Spacer(1, 12))
        
        # Stock Info
        if "stock_info" in data:
            story.append(Paragraph("Stock Information", styles["Heading2"]))
            stock_info = data["stock_info"]
            info_data = [[k.replace("_", " ").title(), str(v)] for k, v in stock_info.items()]
            info_table = Table(info_data, colWidths=[200, 300])
            info_table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (0, -1), colors.grey),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
                ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, 0), 12),
                ("BOTTOMPADDING", (0, 0), (-1, 0), 12),
                ("BACKGROUND", (1, 0), (1, -1), colors.beige),
                ("GRID", (0, 0), (-1, -1), 1, colors.black),
            ]))
            story.append(info_table)
            story.append(Spacer(1, 12))
        
        # Build PDF
        doc.build(story)
        
        return {
            "status": "success",
            "format": "pdf",
            "output_path": output_path,
            "file_size": os.path.getsize(output_path) if os.path.exists(output_path) else 0,
        }
    except Exception as e:
        return {"error": f"Failed to generate PDF report: {str(e)}"}


def _label(key: str) -> str:
    """Title-case snake_case keys; leave human-written labels (e.g. "iPhone Revenue") as-is."""
    return key.replace("_", " ").title() if "_" in key or key.islower() else key


@tool
def generate_markdown_report(data: Dict[str, Any], output_path: str) -> Dict[str, Any]:
    """Generate a Markdown report from research data.
    
    Args:
        data: Dictionary containing research data
        output_path: Path where the Markdown file should be saved
    
    Returns:
        Dictionary with status and file path
    """
    try:
        # Ensure output directory exists
        output_dir = os.path.dirname(output_path)
        if output_dir:
            os.makedirs(output_dir, exist_ok=True)
        
        lines = []
        lines.append("# Financial Research Report\n")
        
        # Metadata
        if "metadata" in data:
            metadata = data["metadata"]
            if metadata.get("company_name"):
                lines.append(f"**Company:** {metadata['company_name']}")
            lines.append(f"**Ticker:** {metadata.get('ticker', 'N/A')}")
            lines.append(f"**Generated At:** {metadata.get('generated_at', 'N/A')}\n")
        
        # Stock Info
        if "stock_info" in data:
            lines.append("## Stock Information\n")
            stock_info = data["stock_info"]
            for key, value in stock_info.items():
                lines.append(f"- **{_label(key)}:** {value}")
            lines.append("")
        
        # Metrics
        if "metrics" in data:
            lines.append("## Financial Metrics\n")
            metrics = data["metrics"]
            for key, value in metrics.items():
                lines.append(f"- **{_label(key)}:** {value}")
            lines.append("")
        
        # Analysis
        if "analysis" in data:
            lines.append("## Analysis\n")
            analysis = data["analysis"]
            if isinstance(analysis, dict):
                for key, value in analysis.items():
                    lines.append(f"### {_label(key)}\n")
                    if isinstance(value, dict):
                        for k, v in value.items():
                            lines.append(f"- **{_label(k)}:** {v}")
                    elif isinstance(value, list):
                        lines.extend(f"- {item}" for item in value)
                    else:
                        lines.append(str(value))
                    lines.append("")
        
        # Write to file
        with open(output_path, "w", encoding="utf-8") as f:
            f.write("\n".join(lines))
        
        return {
            "status": "success",
            "format": "markdown",
            "output_path": output_path,
            "file_size": os.path.getsize(output_path) if os.path.exists(output_path) else 0,
        }
    except Exception as e:
        return {"error": f"Failed to generate Markdown report: {str(e)}"}


@tool
def generate_json_report(data: Dict[str, Any], output_path: str) -> Dict[str, Any]:
    """Generate a JSON report from research data.
    
    Args:
        data: Dictionary containing research data
        output_path: Path where the JSON file should be saved
    
    Returns:
        Dictionary with status and file path
    """
    try:
        # Ensure output directory exists
        output_dir = os.path.dirname(output_path)
        if output_dir:
            os.makedirs(output_dir, exist_ok=True)
        
        # Add metadata if not present
        if "metadata" not in data:
            ticker = data.get("ticker", data.get("stock_info", {}).get("ticker", "UNKNOWN"))
            data["metadata"] = create_report_metadata(ticker)
        
        # Write JSON file
        with open(output_path, "w", encoding="utf-8") as f:
            json.dump(data, f, indent=2, default=str)
        
        return {
            "status": "success",
            "format": "json",
            "output_path": output_path,
            "file_size": os.path.getsize(output_path) if os.path.exists(output_path) else 0,
        }
    except Exception as e:
        return {"error": f"Failed to generate JSON report: {str(e)}"}


# Export all tools
REPORTING_TOOLS = [
    generate_excel_report,
    generate_pdf_report,
    generate_markdown_report,
    generate_json_report,
]
